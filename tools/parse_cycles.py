"""Парсер «РАСПИСАНИЕ ЦИКЛОВЫХ ЗАНЯТИЙ» (4-6 курс, датированные циклы) -> JSON."""
import json
import os
import re
import sys
import warnings

import openpyxl

warnings.filterwarnings("ignore")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAW = os.path.join(ROOT, "data", "raw")
OUT = os.path.join(ROOT, "data", "parsed_cycles")


def cell(g, r, c):
    v = g.get((r, c))
    if v is None:
        return ""
    if hasattr(v, "year"):
        return v.strftime("%d.%m.%Y")
    if hasattr(v, "hour"):
        return v.strftime("%H:%M")
    return str(v).strip()


def fmt_date(v):
    if hasattr(v, "strftime"):
        return v.strftime("%d.%m.%Y")
    s = str(v or "").strip()
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})(?:[ T].*)?$", s)
    if m:
        return f"{m.group(3)}.{m.group(2)}.{m.group(1)}"
    return s


def parse(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = max(wb.worksheets, key=lambda w: w.max_row)
    grid = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value not in (None, ""):
                grid[(c.row, c.column)] = c.value
    for rng in ws.merged_cells.ranges:
        v = grid.get((rng.min_row, rng.min_col))
        if v in (None, ""):
            continue
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                grid.setdefault((r, c), v)

    meta = {"direction": "", "course": None}
    head_row = None
    for r in range(1, 30):
        line = " ".join(cell(grid, r, c) for c in range(1, 5)).lower()
        if "направление" in line:
            meta["direction"] = " ".join(cell(grid, r, c) for c in range(1, 6)).split(":", 1)[-1].strip()
        m = re.search(r"(\d)\s*курс", line)
        if m and not meta["course"]:
            meta["course"] = int(m.group(1))
        if cell(grid, r, 2).lower().startswith("группа"):
            head_row = r
            break
    if not head_row:
        return None

    cycles = []
    by_key = {}
    last_group = ""
    for r in range(head_row + 1, ws.max_row + 1):
        grp = cell(grid, r, 2)
        subj = cell(grid, r, 3)
        teacher = cell(grid, r, 4)
        loc = cell(grid, r, 5)
        room = cell(grid, r, 6)
        dates = cell(grid, r, 8)
        hours = cell(grid, r, 9)
        start = cell(grid, r, 10)[:5]
        end = cell(grid, r, 11)[:5]
        if not any([grp, subj, dates, hours, start]):
            continue
        if grp:
            last_group = grp
        if not subj or not last_group:
            continue
        key = (last_group, subj)
        cur = by_key.get(key)
        if cur is None:
            cur = {"group": last_group, "subject": subj, "teacher": teacher,
                   "place": ", ".join(x for x in (loc, room) if x), "periods": []}
            by_key[key] = cur
            cycles.append(cur)
        if dates or hours or start:
            p = {"dates": dates, "hours": hours, "start": start, "end": end}
            if p not in cur["periods"]:
                cur["periods"].append(p)
    return {"meta": meta, "cycles": [c for c in cycles if c["periods"]]}


def main():
    os.makedirs(OUT, exist_ok=True)
    man = json.load(open(os.path.join(ROOT, "data", "manifest.json"), encoding="utf-8"))
    n = 0
    for m in man:
        if "cikly" not in m["key"]:
            continue
        data = parse(os.path.join(RAW, m["file"]))
        if not data:
            print("FAIL", m["key"])
            continue
        data["key"] = m["key"]
        data["label"] = m.get("label", "")
        json.dump(data, open(os.path.join(OUT, m["key"] + ".json"), "w", encoding="utf-8"),
                  ensure_ascii=False, indent=1)
        n += 1
        print(m["key"], "циклов:", len(data["cycles"]), "групп:",
              len({c["group"] for c in data["cycles"]}))
    print("готово:", n)


if __name__ == "__main__":
    sys.exit(main())
