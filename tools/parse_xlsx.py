"""Парсер Excel-расписаний КГМУ (недельная сетка) -> JSON."""
import json
import os
import re
import sys
import warnings

import openpyxl

warnings.filterwarnings("ignore")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAW = os.path.join(ROOT, "data", "raw")
OUT = os.path.join(ROOT, "data", "parsed")
DAYS = {"ПН": 1, "ВТ": 2, "СР": 3, "ЧТ": 4, "ПТ": 5, "СБ": 6, "ВС": 7}
TYPES = {"лекция": "lecture", "практика": "practice", "цикл": "cycle", "семинар": "seminar"}


def split_top(s, sep=" / "):
    """Делит строку по разделителю, не задевая текст внутри скобок."""
    out, buf, depth = [], [], 0
    i = 0
    while i < len(s):
        ch = s[i]
        if ch in "([":
            depth += 1
        elif ch in ")]":
            depth = max(0, depth - 1)
        if depth == 0 and s.startswith(sep, i):
            out.append("".join(buf))
            buf = []
            i += len(sep)
            continue
        buf.append(ch)
        i += 1
    out.append("".join(buf))
    return [x.strip() for x in out if x.strip()]


def parse_cell(text):
    """'Лекция: Биохимия (2-11) Мустафин И.Г. (УЛК №2, ауд. 331)' -> dict."""
    t = " ".join(str(text).split())
    if not t:
        return None
    kind = None
    m = re.match(r"^(Лекция|Практика|Цикл|Семинар)\s*:\s*", t, re.I)
    if m:
        kind = TYPES[m.group(1).lower()]
        t = t[m.end():]
    # последние скобки — место проведения
    place = ""
    m = re.match(r"^(.*)\(([^()]*)\)\s*$", t)
    if m:
        place = m.group(2).strip()
        t = m.group(1).strip()
    # преподаватель — хвост после «Фамилия И.О.» или «Не назначено»
    teacher = ""
    m = re.match(r"^(.*?)\s+((?:[А-ЯЁA-Z][а-яёa-z\-]+\s+[А-ЯЁA-Z]\.\s*(?:[А-ЯЁA-Z]\.)?)|(?:Не назначено))$", t)
    if m:
        teacher = m.group(2).strip()
        t = m.group(1).strip()
    # предпоследние скобки — подгруппы
    subgroups = ""
    m = re.match(r"^(.*)\(([\d,\s\-–—]*)\)\s*$", t)
    if m:
        subgroups = m.group(2).strip()
        t = m.group(1).strip()
    return {"type": kind, "subject": t, "subgroups": subgroups, "teacher": teacher, "place": place}


def read_grid(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    best = None
    for ws in wb.worksheets:
        if ws.max_row >= 20 and (best is None or ws.max_row > best.max_row):
            best = ws
    ws = best
    grid = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value not in (None, ""):
                grid[(c.row, c.column)] = c.value
    for rng in ws.merged_cells.ranges:
        v = grid.get((rng.min_row, rng.min_col))
        if v in (None, ""):
            continue
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                grid.setdefault((r, c), v)
    return ws, grid


def cell(g, r, c):
    v = g.get((r, c))
    return "" if v is None else str(v).strip()


def parse_week_file(path):
    ws, g = read_grid(path)
    maxr, maxc = ws.max_row, ws.max_column

    meta = {"direction": "", "level": "", "form": "", "course": None}
    group_row = potok_row = None
    for r in range(1, min(25, maxr) + 1):
        for c in range(1, min(4, maxc) + 1):
            v = cell(g, r, c).lower()
            if v.startswith("направление"):
                meta["direction"] = cell(g, r, c + 2) or cell(g, r, c + 1)
            elif v.startswith("уровень"):
                meta["level"] = cell(g, r, c + 2) or cell(g, r, c + 1)
            elif v.startswith("форма обучения"):
                meta["form"] = cell(g, r, c + 2) or cell(g, r, c + 1)
            elif v == "поиск группы":
                group_row = r
            elif v == "поток" and potok_row is None:
                potok_row = r
        a, b = cell(g, r, 2), cell(g, r, 3).lower()
        if b.startswith("курс") and a.strip().isdigit():
            meta["course"] = int(a)

    if not group_row:
        return None
    groups = []
    for c in range(4, maxc + 1):
        name = cell(g, group_row, c)
        if not name:
            continue
        potok = cell(g, potok_row, c) if potok_row else ""
        groups.append({"name": name, "potok": potok})
    if not groups:
        return None

    data_start = None
    for r in range(max(1, group_row + 1), maxr + 1):
        if cell(g, r, 1).upper() in DAYS:
            data_start = r
            break
    if not data_start:
        return None

    days = {}
    for r in range(data_start, maxr + 1):
        day = cell(g, r, 1).upper()
        if day not in DAYS:
            continue
        start, end = cell(g, r, 2), cell(g, r, 3)
        if not start:
            continue
        start = start[:5]
        end = end[:5] if end else ""
        for i, grp in enumerate(groups):
            raw = cell(g, r, 4 + i)
            if not raw:
                continue
            segs = split_top(raw)
            parsed = [p for p in (parse_cell(s) for s in segs) if p]
            if not parsed:
                continue
            lst = days.setdefault(day, [])
            for p in parsed:
                lst.append(dict(p, group=grp["name"], start=start, end=end, row=r))

    # склеиваем одинаковые занятия в соседних 45-минутных строках в одну пару (по каждой группе)
    day_list = []
    for day in ["ПН", "ВТ", "СР", "ЧТ", "ПТ", "СБ", "ВС"]:
        if day not in days:
            continue
        merged = []
        for grp in groups:
            items = [x for x in days[day] if x["group"] == grp["name"]]
            by_sig = {}
            for it in items:
                sig = (it["type"], it["subject"], it["subgroups"], it["teacher"], it["place"])
                by_sig.setdefault(sig, []).append(it)
            for sig, lst in by_sig.items():
                lst.sort(key=lambda x: x["row"])
                cur = None
                for it in lst:
                    if cur is not None and cur["row"] + 1 == it["row"]:
                        cur["end"] = it["end"]
                        cur["row"] = it["row"]
                        continue
                    cur = dict(it)
                    merged.append(cur)
        for m in merged:
            m.pop("_sig", None)
            m.pop("row", None)
        day_list.append({
            "day": day,
            "wd": DAYS[day],
            "lessons": sorted(merged, key=lambda x: (x["start"], x["group"])),
        })

    return {
        "meta": meta,
        "groups": groups,
        "days": day_list,
        "lessons": sum(len(d["lessons"]) for d in day_list),
    }


def main():
    os.makedirs(OUT, exist_ok=True)
    man = json.load(open(os.path.join(ROOT, "data", "manifest.json"), encoding="utf-8"))
    report, bad = [], []
    for m in man:
        if not m["file"].endswith(".xlsx"):
            continue
        path = os.path.join(RAW, m["file"])
        if "cikly" in m["key"]:
            continue  # цикловые — отдельный формат, позже
        try:
            data = parse_week_file(path)
        except Exception as e:  # noqa: BLE001
            bad.append((m["key"], f"{type(e).__name__}: {e}"))
            continue
        if not data:
            bad.append((m["key"], "не распознана сетка"))
            continue
        data["key"] = m["key"]
        data["program_page"] = m["program"]
        data["label"] = m["label"]
        json.dump(data, open(os.path.join(OUT, m["key"] + ".json"), "w", encoding="utf-8"),
                  ensure_ascii=False, indent=1)
        report.append((m["key"], len(data["groups"]), data["lessons"], len(data["days"])))
    for r in report:
        print("ok  ", r)
    for b in bad:
        print("FAIL", b)
    print(f"готово: {len(report)} файлов, ошибок {len(bad)}")


if __name__ == "__main__":
    sys.exit(main())
